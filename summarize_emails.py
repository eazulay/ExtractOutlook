"""
Iterate JSON files in extracted_emails/ and send each to claude-haiku-4-5
to produce a markdown summary table. Per-file markdown goes to email_summaries/,
and the aggregated result is written to claude_summary.xlsx.

Usage:
    python summarize_emails.py [--input-dir <dir>] [--output-dir <dir>] [--xlsx <path>]
"""

import argparse
import json
import os
import sys
from pathlib import Path

import anthropic
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

load_dotenv()

# Chunk JSON content so a single request stays well under the 50K input
# tokens/min cap on Haiku 4.5 Tier 1. ~3.5 chars/token for JSON, plus
# headroom for the system prompt and rolling-window concurrency.
CHUNK_CHAR_BUDGET = 120_000


PROMPT = """You MUST incorporate information from the attachment_content field into the Summary for every email. If attachment_content is non-empty, its content is the primary source of information for the Summary, even if plain_body is empty.

The attached JSON file contains a collection of emails. Each email object has these fields:
- subject – the email subject line
- sender_name – the name of the sender
- sender_email – the sender's name and email address in the format: Name <email>
- to – the recipient(s) of the email
- cc – any CC recipients (may be empty)
- delivery_time – the date and time the email was sent, output as-is without timezone conversion
- plain_body – the plain text body of the email
- attachment_content – the extracted text of all attachments, if any; each attachment is preceded by its filename in the format === filename ===

Process every email in the file and output one row per email in a markdown table with these exact columns:

| DateTime | From | To | Subject | Summary |

Column definitions:
- DateTime – delivery_time formatted as YYYY-MM-DD HH:MM; use a space between date and time, not the letter T; example: 2024-03-15 09:30
- From – the sender_email field, verbatim
- To – the to field, verbatim; if multiple recipients, separate with semicolons; do not include CC recipients in this column
- Subject – the subject field, with repeated Re: and Fwd: prefixes stripped; otherwise verbatim
- Summary – bullet points combining plain_body and attachment_content; the Summary MUST reflect key details from attachment_content if it is non-empty, even if plain_body is empty; only include attachment content that is substantive and relevant to the email's purpose — ignore boilerplate, legal footers, and standard disclaimers; if attachment content is lengthy, summarise only the key points most relevant to the email's purpose; each bullet covers one key point, decision, request, deadline, or action item; use telegraphic style with no unnecessary words; indicate which attachment a point comes from using its filename; if CC recipients are present, add a final bullet "CC: [names]"; maximum 300 words per email

Rules:
- Output a valid markdown table with a header row and a separator row.
- Output field values only, never the surrounding JSON quotation marks. For example, if sender_email is "\\"John Smith\\" <john@example.com>", output John Smith <john@example.com> with no leading or trailing quotation marks.
- If a field is missing or empty, output a hyphen (-) for that field.
- IMPORTANT: Never output a hyphen (-) for Summary if either plain_body or attachment_content contains text.
- If an email address contains an opening < but no closing >, add the closing > immediately after the email address.
- Summarise only the new content at the top of plain_body; ignore quoted reply chains beginning with > or From: headers.
- Do not add any text, commentary, or blank lines before or after the table.
- In the Summary column, use * for bullets and <br> to separate them; example: * First point<br>* Second point<br>* Third point

Generate a markdown table with the requested columns. Do not include any commentary or explanations outside the table."""


def chunk_emails(emails: list, char_budget: int = CHUNK_CHAR_BUDGET) -> list[list]:
    """Split a list of email objects into groups whose serialized size fits the budget."""
    chunks: list[list] = []
    current: list = []
    current_size = 2  # JSON array brackets
    for email in emails:
        size = len(json.dumps(email, ensure_ascii=False)) + 2  # +2 for ", "
        if current and current_size + size > char_budget:
            chunks.append(current)
            current = []
            current_size = 2
        current.append(email)
        current_size += size
    if current:
        chunks.append(current)
    return chunks


def strip_table_header(md: str) -> str:
    """Drop the first two pipe-rows (header + separator) from a markdown table."""
    lines = md.splitlines()
    pipe_count = 0
    for i, line in enumerate(lines):
        if line.strip().startswith("|"):
            pipe_count += 1
            if pipe_count == 2:
                return "\n".join(lines[i + 1 :])
    return md


def summarize_chunk(client: anthropic.Anthropic, payload: str, label: str) -> tuple[str, anthropic.types.Usage]:
    with client.messages.stream(
        model="claude-haiku-4-5",
        max_tokens=64000,
        system=[
            {
                "type": "text",
                "text": PROMPT,
                "cache_control": {"type": "ephemeral"},
            }
        ],
        messages=[
            {
                "role": "user",
                "content": f"Here is the JSON file content:\n\n{payload}",
            }
        ],
    ) as stream:
        parts = []
        for text in stream.text_stream:
            parts.append(text)
            print(text, end="", flush=True)
        final = stream.get_final_message()
    usage = final.usage
    print(
        f"\n  -> {label} "
        f"[input={usage.input_tokens} cache_read={usage.cache_read_input_tokens} "
        f"cache_write={usage.cache_creation_input_tokens} output={usage.output_tokens}]"
    )
    return "".join(parts), usage


def summarize_file(client: anthropic.Anthropic, json_path: Path, output_path: Path) -> None:
    raw = json_path.read_text(encoding="utf-8")

    # If it fits in one shot, send as-is so we don't pay the parse/reserialize cost.
    if len(raw) <= CHUNK_CHAR_BUDGET:
        md, _ = summarize_chunk(client, raw, output_path.name)
        output_path.write_text(md, encoding="utf-8")
        return

    emails = json.loads(raw)
    if not isinstance(emails, list):
        # Not the expected shape — fall back to single request and let Claude or the API complain.
        md, _ = summarize_chunk(client, raw, output_path.name)
        output_path.write_text(md, encoding="utf-8")
        return

    groups = chunk_emails(emails)
    print(f"  Splitting {len(emails)} emails into {len(groups)} chunk(s) (file too large for one request)")

    combined_parts: list[str] = []
    for idx, group in enumerate(groups, 1):
        payload = json.dumps(group, ensure_ascii=False, indent=2)
        label = f"{output_path.name} chunk {idx}/{len(groups)}"
        print(f"\n  -- chunk {idx}/{len(groups)} ({len(group)} emails, {len(payload):,} chars)")
        md, _ = summarize_chunk(client, payload, label)
        combined_parts.append(md if idx == 1 else strip_table_header(md))

    output_path.write_text("\n".join(combined_parts), encoding="utf-8")


def folder_name_from_json(json_stem: str) -> str:
    """Extract the folder path portion of a json filename.

    "GIW up to 2004_GIW.Others.Expensive.Abacus FS" -> "GIW.Others.Expensive.Abacus FS"
    "GIW up to 2004_2004 Tax Claim" -> "2004 Tax Claim"
    """
    # Drop the leading "<archive name>_" prefix; keep the full period-delimited path.
    if "_" in json_stem:
        json_stem = json_stem.split("_", 1)[1]
    return json_stem


def parse_markdown_table(md: str) -> list[list[str]]:
    """Return data rows (no header, no separator) from a markdown pipe table."""
    rows = []
    for line in md.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if not cells:
            continue
        # Skip the separator row (--- | --- | ...)
        if all(set(c) <= set("-:") and c for c in cells):
            continue
        rows.append(cells)
    # First remaining row is the header
    return rows[1:] if rows else []


def write_xlsx(md_dir: Path, json_files: list[Path], xlsx_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summaries"

    headers = ["Folder", "DateTime", "From", "To", "Subject", "Summary"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    total_rows = 0
    for json_path in json_files:
        md_path = md_dir / f"{json_path.stem}.md"
        if not md_path.exists():
            print(f"  (skipping xlsx row — no markdown for {json_path.name})")
            continue
        folder = folder_name_from_json(json_path.stem)
        for row in parse_markdown_table(md_path.read_text(encoding="utf-8")):
            # Pad or truncate to 5 columns to match the header
            row = (row + ["-"] * 5)[:5]
            # Convert <br> in the Summary column to real newlines
            row[4] = row[4].replace("<br>", "\n")
            ws.append([folder, *row])
            total_rows += 1

    widths = {"A": 22, "B": 18, "C": 36, "D": 36, "E": 50, "F": 90}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap

    ws.freeze_panes = "A2"
    wb.save(xlsx_path)
    print(f"\nWrote {total_rows} row(s) to {xlsx_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Summarize emails with Claude Haiku 4.5.")
    parser.add_argument("--input-dir", default="extracted_emails", help="Directory containing JSON files")
    parser.add_argument("--output-dir", default="email_summaries", help="Where to write markdown summaries")
    parser.add_argument("--xlsx", default="claude_summary.xlsx", help="Aggregated Excel output path")
    parser.add_argument("--overwrite", action="store_true", help="Reprocess files that already have output")
    parser.add_argument("--xlsx-only", action="store_true", help="Skip Claude calls; just rebuild the xlsx from existing markdown")
    args = parser.parse_args()

    input_dir = Path(args.input_dir)
    output_dir = Path(args.output_dir)

    if not input_dir.is_dir():
        print(f"Error: {input_dir} is not a directory")
        sys.exit(1)

    output_dir.mkdir(parents=True, exist_ok=True)

    json_files = sorted(input_dir.glob("*.json"))
    if not json_files:
        print(f"No JSON files found in {input_dir}")
        sys.exit(1)

    print(f"Found {len(json_files)} JSON file(s) in {input_dir}\n")

    if not args.xlsx_only:
        if not os.environ.get("ANTHROPIC_API_KEY"):
            print("Error: ANTHROPIC_API_KEY is not set")
            sys.exit(1)

        # max_retries=10 with exponential backoff lets the SDK ride out
        # rate-limit windows (honours the retry-after header on 429s).
        client = anthropic.Anthropic(max_retries=10, timeout=600.0)

        for i, json_path in enumerate(json_files, 1):
            output_path = output_dir / f"{json_path.stem}.md"
            if output_path.exists() and not args.overwrite:
                print(f"[{i}/{len(json_files)}] {json_path.name} -> already exists, skipping")
                continue

            print(f"[{i}/{len(json_files)}] {json_path.name}")
            try:
                summarize_file(client, json_path, output_path)
            except anthropic.APIError as e:
                print(f"\n  ERROR: {e}")
            print()

    write_xlsx(output_dir, json_files, Path(args.xlsx))


if __name__ == "__main__":
    main()
